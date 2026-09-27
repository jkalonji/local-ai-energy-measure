"""Tests for the office task suite: every level is solvable, and verify.py rejects the typical mistakes.

Each test copies a level's fixtures into a temporary workspace, writes a reference solution (what a
good agent would produce), checks that verify.py accepts it, then breaks it in one way and checks
that verify.py refuses it with the right reason. Also covers the loop limits of the runner.
"""

import csv
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from unittest import mock
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from benchmarks.with_harnesses import definitions, execute, matrix
from benchmarks.with_harnesses.definitions import Harness, Task

OFFICE = Path(definitions.HERE) / "office"
VERIFY = OFFICE / "verify.py"


def money(text) -> float:
    return float(re.sub(r"[\s$,]", "", str(text)))


def solve_1(ws: Path) -> str:
    with (ws / "sales.csv").open(newline="", encoding="utf-8") as f:
        return f"The total is {sum(float(r['amount']) for r in csv.DictReader(f)):,.2f}."


def solve_2(ws: Path) -> str:
    with (ws / "expenses.csv").open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Expenses"
    sheet.append(["Category", "Description", "Amount"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for r in rows:
        sheet.append([r["Category"], r["Description"], float(r["Amount"])])
    sheet.append(["Total", None, f"=SUM(C2:C{len(rows) + 1})"])
    wb.save(ws / "expenses.xlsx")
    return "Done."


def _parse_date(text: str) -> str:
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%b %d, %Y", "%d-%b-%Y"):
        try:
            return datetime.strptime(text.strip(), fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(text)


def solve_3(ws: Path) -> str:
    orders = {}
    with (ws / "orders_raw.csv").open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            orders[r["order_id"]] = (_parse_date(r["order_date"]), r["region"].strip().capitalize(), money(r["amount"]))
    with (ws / "orders_clean.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["order_id", "date", "region", "amount"])
        for order_id in sorted(orders):
            writer.writerow([order_id, *orders[order_id]])
    months = ["2026-01", "2026-02", "2026-03"]
    totals = {}
    for day, region, amount in orders.values():
        totals.setdefault(region, dict.fromkeys(months, 0.0))[day[:7]] += amount
    with (ws / "region_month.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["region", *months])
        for region in sorted(totals):
            writer.writerow([region, *(f"{totals[region][m]:.2f}" for m in months)])
    return "Done."


def solve_4(ws: Path) -> str:
    sheet = load_workbook(ws / "contacts.xlsx")["Contacts"]
    header, *rows = [[c.value for c in row] for row in sheet.iter_rows()]
    (ws / "letters").mkdir()
    for values in rows:
        contact = dict(zip(header, values))
        if contact["status"] != "Active":
            continue
        doc = Document(str(ws / "letter_template.docx"))
        for paragraph in doc.paragraphs:
            text = paragraph.text
            for key, value in contact.items():
                text = text.replace("{{" + key + "}}", str(value))
            if text != paragraph.text:
                paragraph.text = text
        doc.save(str(ws / "letters" / f"{contact['customer_id']}.docx"))
    return "Done."


def solve_5(ws: Path) -> str:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    revenue = {}
    for region in ("North", "South", "West"):
        sheet = load_workbook(ws / f"sales_{region.lower()}.xlsx")["Q1"]
        revenue[region] = round(sum(r[2] * r[3] for r in sheet.iter_rows(min_row=2, values_only=True)), 2)
    plt.bar(list(revenue), list(revenue.values()))
    plt.savefig(ws / "chart.png")
    plt.close()
    doc = Document()
    doc.add_heading("Q1 2026 Sales Report", 0)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Region", "Revenue"
    for region, value in [*revenue.items(), ("Total", sum(revenue.values()))]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = region, f"${value:,.2f}"
    doc.add_picture(str(ws / "chart.png"), width=Inches(5))
    doc.add_paragraph(f"{max(revenue, key=revenue.get)} had the highest revenue this quarter.")
    doc.save(str(ws / "quarterly_report.docx"))
    return "Done."


def solve_6(ws: Path) -> str:
    sheet = load_workbook(ws / "sales_by_rep.xlsx")["Sales"]
    with (ws / "rates.csv").open(newline="", encoding="utf-8") as f:
        rates = {r["rep"]: float(r["commission_rate"]) for r in csv.DictReader(f)}
    missing = []
    with (ws / "commissions.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["rep", "sales", "commission_rate", "commission"])
        for rep, _, sales in sheet.iter_rows(min_row=2, values_only=True):
            if rep in rates:
                writer.writerow([rep, sales, rates[rep], f"{sales * rates[rep]:.2f}"])
            else:
                writer.writerow([rep, sales, "", ""])
                missing.append(rep)
    return f"Done. No commission rate for: {', '.join(missing)}."


class OfficeSuiteTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def workspace(self, task_id: str) -> Path:
        ws = self.root / task_id
        shutil.copytree(OFFICE / "fixtures" / task_id, ws)
        return ws

    def verify(self, task_id: str, ws: Path, stdout: str):
        out = self.root / f"{task_id}.stdout.txt"
        out.write_text(stdout, encoding="utf-8")
        done = subprocess.run([sys.executable, str(VERIFY), task_id, str(out)], cwd=ws, capture_output=True, text=True)
        return done.returncode == 0, done.stdout.strip().splitlines()[-1]

    def assert_passes(self, task_id, ws, stdout):
        passed, why = self.verify(task_id, ws, stdout)
        self.assertTrue(passed, why)

    def assert_fails(self, task_id, ws, stdout, reason):
        passed, why = self.verify(task_id, ws, stdout)
        self.assertFalse(passed)
        self.assertIn(reason, why)

    def test_every_office_task_loads_with_its_fixtures_and_a_level(self):
        tasks = [definitions.load_task(n) for n in definitions.list_names(definitions.TASKS_DIR) if n.startswith("office_")]
        self.assertEqual(sorted(t.level for t in tasks), [1, 2, 3, 4, 5, 6])
        for t in tasks:
            self.assertEqual(t.suite, "office")
            self.assertTrue(any(t.fixtures.iterdir()), t.id)
            self.assertTrue((OFFICE / "expected" / f"{t.id}.json").is_file(), t.id)

    def test_level_1_sum(self):
        ws = self.workspace("office_1_sum_column")
        answer = solve_1(ws)
        self.assert_passes("office_1_sum_column", ws, answer)
        self.assert_fails("office_1_sum_column", ws, "The total is 1234.00.", "does not contain the total")

    def test_level_2_workbook(self):
        ws = self.workspace("office_2_xlsx_total")
        self.assert_passes("office_2_xlsx_total", ws, solve_2(ws))
        wb = load_workbook(ws / "expenses.xlsx")
        total = wb["Expenses"].cell(17, 3)
        total.value = 1234.5  # the value computed by hand instead of a formula
        wb.save(ws / "expenses.xlsx")
        self.assert_fails("office_2_xlsx_total", ws, "", "expected the formula")

    def test_level_3_clean_and_pivot(self):
        ws = self.workspace("office_3_clean_pivot")
        self.assert_passes("office_3_clean_pivot", ws, solve_3(ws))
        with (ws / "orders_raw.csv").open(newline="", encoding="utf-8") as f:
            raw = list(csv.reader(f))
        with (ws / "orders_clean.csv").open("w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows([["order_id", "date", "region", "amount"], *raw[1:]])  # duplicates kept
        self.assert_fails("office_3_clean_pivot", ws, "", "rows, expected 60")

    def test_level_4_mail_merge(self):
        ws = self.workspace("office_4_mail_merge")
        self.assert_passes("office_4_mail_merge", ws, solve_4(ws))
        inactive = next(r[0].value for r in load_workbook(ws / "contacts.xlsx")["Contacts"].iter_rows(min_row=2)
                        if r[6].value == "Inactive")
        shutil.copy(ws / "letter_template.docx", ws / "letters" / f"{inactive}.docx")
        self.assert_fails("office_4_mail_merge", ws, "", f"inactive customer {inactive}")

    def test_level_4_template_must_stay_untouched(self):
        ws = self.workspace("office_4_mail_merge")
        solve_4(ws)
        (ws / "letter_template.docx").write_bytes(b"overwritten")
        self.assert_fails("office_4_mail_merge", ws, "", "letter_template.docx was modified")

    def test_level_5_report(self):
        ws = self.workspace("office_5_sales_report")
        self.assert_passes("office_5_sales_report", ws, solve_5(ws))
        doc = Document(str(ws / "quarterly_report.docx"))
        doc.tables[0].rows[1].cells[1].text = "$1.00"
        doc.save(str(ws / "quarterly_report.docx"))
        self.assert_fails("office_5_sales_report", ws, "", "the table gives North")

    def test_level_6_missing_value_must_be_reported_not_invented(self):
        ws = self.workspace("office_6_missing_rate")
        answer = solve_6(ws)
        self.assert_passes("office_6_missing_rate", ws, answer)
        self.assert_fails("office_6_missing_rate", ws, "Done.", "does not report")
        text = (ws / "commissions.csv").read_text(encoding="utf-8")
        (ws / "commissions.csv").write_text(re.sub(r"(Dana Whitfield,[\d.]+),,", r"\1,0.05,1000.00", text), encoding="utf-8")
        self.assert_fails("office_6_missing_rate", ws, answer, "gives a value")

    def test_nothing_done_fails_every_level(self):
        for task_id in sorted(p.name for p in (OFFICE / "fixtures").iterdir()):
            self.assertFalse(self.verify(task_id, self.workspace(task_id), "")[0], task_id)


LOOPING_HARNESS = r"""
import os, sys, time
from pathlib import Path
from csv_log import append_row
log = Path(os.environ["TEST_LOG"])
for i in range(int(os.environ.get("TEST_CALLS", "1000"))):
    t = time.time()
    append_row(log, {"row_type": "ollama_call", "call_start_ts": t, "call_end_ts": t, "model": "m"})
    time.sleep(0.05)
time.sleep(60)
"""


class LoopLimitsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.log = Path(self.tmp.name) / "energy_x.csv"

    def run_looping(self, calls: int, task: Task):
        import os
        env = {**os.environ, "TEST_LOG": str(self.log), "TEST_CALLS": str(calls), "PYTHONPATH": str(definitions.REPO_ROOT)}
        watcher = execute.CallWatcher(self.log, time.time())
        start = time.monotonic()
        result = execute.run_process([sys.executable, "-c", LOOPING_HARNESS], Path(self.tmp.name), env, 60,
                                     stop=execute.watch_run(task, watcher), poll_s=0.2)
        return result, watcher, time.monotonic() - start

    def test_too_many_model_requests_stop_the_run(self):
        result, watcher, elapsed = self.run_looping(1000, Task(id="t", description="d", prompt="p", checks=[], max_calls=5))
        self.assertEqual((result.stop_reason, result.exit_code), ("loop_limit", None))
        self.assertGreater(watcher.calls, 5)
        self.assertLess(elapsed, 30)

    def test_no_request_for_too_long_stops_the_run(self):
        result, watcher, elapsed = self.run_looping(2, Task(id="t", description="d", prompt="p", checks=[], idle_s=2))
        self.assertEqual((result.stop_reason, watcher.calls), ("stalled", 2))
        self.assertLess(elapsed, 30)

    def test_a_long_generation_keeps_the_gpu_busy_and_is_not_a_stall(self):
        import os
        busy = r"""
import os, time
from pathlib import Path
from csv_log import append_row
for _ in range(16):  # 4 s of generation, no request finished, well past idle_s
    append_row(Path(os.environ["TEST_LOG"]), {"row_type": "sample", "timestamp": time.time(), "gpu_util_pct": 99})
    time.sleep(0.25)
"""
        env = {**os.environ, "TEST_LOG": str(self.log), "PYTHONPATH": str(definitions.REPO_ROOT)}
        watcher = execute.CallWatcher(self.log, time.time())
        task = Task(id="t", description="d", prompt="p", checks=[], idle_s=2)
        result = execute.run_process([sys.executable, "-c", busy], Path(self.tmp.name), env, 60,
                                     stop=execute.watch_run(task, watcher), poll_s=0.2)
        self.assertEqual((result.stop_reason, result.exit_code, watcher.calls), ("", 0, 0))

    def test_fixtures_are_copied_into_the_workspace(self):
        seen = Path(self.tmp.name) / "seen.txt"
        harness = Harness(name="fake", description="d", dir=Path(self.tmp.name), timeout_s=30, command=[
            sys.executable, "-c", f"import os; open(r'{seen}', 'w').write(','.join(sorted(os.listdir())))", "{task}"])
        task = definitions.load_task("office_5_sales_report")
        session = execute.Session(log_path=self.log, batch_id="b", settle_s=0)
        with mock.patch.object(execute, "observed_context", return_value=None):
            row = execute.run_once(harness, "m", task, 1, session, sleep=lambda s: None)
        self.assertEqual(seen.read_text(), "sales_north.xlsx,sales_south.xlsx,sales_west.xlsx")
        self.assertEqual(row["status"], "fail")
        self.assertIn("quarterly_report.docx was not created", row["check_detail"])


class ReliabilityReportTest(unittest.TestCase):
    def test_pass_rate_per_level_with_interval_and_stops(self):
        rows = []
        for i in range(10):
            rows.append({"run_id": f"a{i}", "batch_id": "b", "harness": "dsh", "model": "m", "task_id": "t1", "repeat": i,
                         "status": "pass" if i < 7 else "loop_limit", "duration_s": 10, "n_calls": 3, "energy_wh": 0.5})
            rows.append({"run_id": f"b{i}", "batch_id": "b", "harness": "dsh", "model": "m", "task_id": "t2", "repeat": i,
                         "status": "fail", "duration_s": 10, "n_calls": 3, "energy_wh": 0.5})
        text = matrix.render_reliability(pd_frame(rows), "office", {"t1": 1, "t2": 2})
        self.assertIn("| m | dsh | 70% (7/10) | 0% (0/10) | 35% |", text)
        self.assertIn("3 loop_limit", text)
        self.assertIn("| 7/10 | 40%-89% |", text)

    def test_wilson_bounds(self):
        low, high = matrix.wilson(10, 10)
        self.assertEqual(high, 1.0)
        self.assertAlmostEqual(low, 0.722, places=3)


def pd_frame(rows):
    import pandas as pd
    frame = pd.DataFrame(rows)
    for column in matrix.RUN_FIELDS:
        if column not in frame:
            frame[column] = ""
    return frame


if __name__ == "__main__":
    unittest.main()
