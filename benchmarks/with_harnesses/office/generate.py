"""Generate the input files and the expected results of the office task suite.

Every office task (tasks/office_*.yaml) starts from a folder of input files, copied into the run's
workspace so the agent discovers them there. This script writes them, reproducibly (fixed seed):

    office/fixtures/<task_id>/...     the files the agent sees
    office/expected/<task_id>.json    the ground truth, read by verify.py only (never copied)

The generated files are committed; re-run this script only to change the data:
    python benchmarks/with_harnesses/office/generate.py
"""

import csv
import json
import random
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, List

from docx import Document
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
FIXTURES_DIR = HERE / "fixtures"
EXPECTED_DIR = HERE / "expected"
SEED = 20260927
FIXED_TIME = datetime(2026, 1, 1)  # document metadata, so re-generating gives the same content

CUSTOMERS = ["Acme Corp", "Globex", "Initech", "Umbrella Ltd", "Stark Industries", "Wayne Enterprises",
             "Hooli", "Vandelay Imports", "Soylent Co", "Tyrell Corp"]
PRODUCTS = {"Printer paper (box)": 24.90, "Stapler": 12.50, "Desk lamp": 39.00, "Toner cartridge": 89.99,
            "Notebook pack": 8.75, "Office chair": 179.00, "Whiteboard markers": 6.40, "Monitor stand": 45.50}


def _money(value: float) -> float:
    return round(value + 1e-9, 2)


def _write_csv(path: Path, header: List[str], rows: List[list]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)


def _save_workbook(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.properties.created = workbook.properties.modified = FIXED_TIME
    workbook.save(path)


def level_1(rng: random.Random, out: Path) -> Dict:
    """sales.csv: 40 order lines; the agent must report the total of the amount column."""
    rows, start = [], date(2026, 3, 2)
    for i in range(40):
        product = rng.choice(sorted(PRODUCTS))
        quantity = rng.randint(1, 12)
        rows.append([f"SO-{1001 + i}", (start + timedelta(days=rng.randint(0, 27))).isoformat(),
                     rng.choice(CUSTOMERS), product, quantity, f"{PRODUCTS[product]:.2f}",
                     f"{_money(quantity * PRODUCTS[product]):.2f}"])
    _write_csv(out / "sales.csv", ["order_id", "date", "customer", "product", "quantity", "unit_price", "amount"], rows)
    return {"total": _money(sum(float(r[6]) for r in rows))}


def level_2(rng: random.Random, out: Path) -> Dict:
    """expenses.csv: 15 expense lines to turn into a formatted workbook with a SUM formula."""
    items = [("Travel", "Train ticket to Lyon"), ("Travel", "Taxi to the airport"), ("Meals", "Team lunch"),
             ("Office", "Printer ink"), ("Software", "Diagram tool licence"), ("Meals", "Client dinner"),
             ("Travel", "Hotel, two nights"), ("Office", "Ergonomic keyboard"), ("Training", "Online course"),
             ("Software", "Cloud storage plan"), ("Office", "Desk organiser"), ("Meals", "Coffee for workshop"),
             ("Travel", "Parking fee"), ("Training", "Conference ticket"), ("Office", "USB-C hub")]
    rows = [[category, description, f"{_money(rng.uniform(8, 450)):.2f}"] for category, description in items]
    _write_csv(out / "expenses.csv", ["Category", "Description", "Amount"], rows)
    return {"rows": [[c, d, float(a)] for c, d, a in rows], "total": _money(sum(float(r[2]) for r in rows))}


def _messy_date(day: date, rng: random.Random) -> str:
    style = rng.randrange(4)
    if style == 0:
        return day.isoformat()
    if style == 1:
        return day.strftime("%m/%d/%Y")               # US format, as the prompt states
    if style == 2:
        return f"{day.strftime('%b')} {day.day}, {day.year}"
    return day.strftime("%d-%b-%Y")


def _messy_region(region: str, rng: random.Random) -> str:
    return rng.choice([region, region.lower(), region.upper(), f" {region} ", f"{region}  "])


def _messy_amount(amount: float, rng: random.Random) -> str:
    return rng.choice([f"{amount:.2f}", f"${amount:,.2f}", f"{amount:,.2f}", f"{amount:g}"])


def level_3(rng: random.Random, out: Path) -> Dict:
    """orders_raw.csv: mixed date formats, messy region names and amounts, duplicated orders."""
    regions = ["North", "South", "East", "West"]
    orders = []
    for i in range(60):
        day = date(2026, 1, 1) + timedelta(days=rng.randint(0, 89))
        orders.append({"order_id": f"ORD-{2001 + i}", "date": day, "region": rng.choice(regions),
                       "amount": _money(rng.uniform(40, 2500))})
    raw = [[o["order_id"], _messy_date(o["date"], rng), _messy_region(o["region"], rng),
            _messy_amount(o["amount"], rng)] for o in orders]
    for o in rng.sample(orders, 8):  # the same order exported twice, formatted differently the second time
        raw.append([o["order_id"], _messy_date(o["date"], rng), _messy_region(o["region"], rng),
                    _messy_amount(o["amount"], rng)])
    rng.shuffle(raw)
    _write_csv(out / "orders_raw.csv", ["order_id", "order_date", "region", "amount"], raw)

    months = ["2026-01", "2026-02", "2026-03"]
    pivot = {r: {m: 0.0 for m in months} for r in regions}
    for o in orders:
        pivot[o["region"]][o["date"].strftime("%Y-%m")] += o["amount"]
    return {
        "clean": [[o["order_id"], o["date"].isoformat(), o["region"], o["amount"]]
                  for o in sorted(orders, key=lambda o: o["order_id"])],
        "months": months,
        "pivot": {r: {m: _money(v) for m, v in by_month.items()} for r, by_month in sorted(pivot.items())},
    }


def level_4(rng: random.Random, out: Path) -> Dict:
    """letter_template.docx + contacts.xlsx: one personalised letter per active customer."""
    doc = Document()
    doc.add_paragraph("Northwind Office Supplies")
    doc.add_paragraph("Dear {{first_name}},")
    doc.add_paragraph("Thank you for your continued trust. The subscription of {{company}} renews on "
                      "{{renewal_date}} for a yearly amount of {{amount}}.")
    doc.add_paragraph("If you have any question, simply reply to this letter.")
    doc.add_paragraph("Best regards,")
    doc.add_paragraph("The Customer Success Team")
    doc.core_properties.created = doc.core_properties.modified = FIXED_TIME
    out.mkdir(parents=True, exist_ok=True)
    doc.save(out / "letter_template.docx")

    people = [("Alice", "Martin"), ("Bruno", "Keller"), ("Chloe", "Dubois"), ("Daniel", "Ortiz"), ("Emma", "Novak"),
              ("Farid", "Haddad"), ("Grace", "Okafor"), ("Hugo", "Lindqvist"), ("Ines", "Moreau"), ("Jonas", "Weber")]
    inactive = set(rng.sample(range(len(people)), 3))
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Contacts"
    sheet.append(["customer_id", "first_name", "last_name", "company", "renewal_date", "amount", "status"])
    contacts = []
    for i, (first, last) in enumerate(people):
        row = {"customer_id": f"C{101 + i}", "first_name": first, "last_name": last, "company": CUSTOMERS[i],
               "renewal_date": (date(2026, 10, 1) + timedelta(days=rng.randint(0, 90))).strftime("%B %d, %Y"),
               "amount": f"${rng.randrange(600, 9000, 50):,.2f}", "status": "Inactive" if i in inactive else "Active"}
        sheet.append(list(row.values()))
        contacts.append(row)
    _save_workbook(wb, out / "contacts.xlsx")
    return {"letters": {c["customer_id"]: {k: c[k] for k in ("first_name", "company", "renewal_date", "amount")}
                        for c in contacts if c["status"] == "Active"},
            "inactive": sorted(c["customer_id"] for c in contacts if c["status"] == "Inactive")}


def level_5(rng: random.Random, out: Path) -> Dict:
    """sales_<region>.xlsx for three regions: consolidate into a Word report with a table and a chart."""
    revenue = {}
    for region in ("North", "South", "West"):
        wb = Workbook()
        sheet = wb.active
        sheet.title = "Q1"
        sheet.append(["date", "product", "units", "unit_price"])
        total = 0.0
        for _ in range(rng.randint(20, 30)):
            product = rng.choice(sorted(PRODUCTS))
            units = rng.randint(1, 40)
            day = date(2026, 1, 1) + timedelta(days=rng.randint(0, 89))
            sheet.append([day.isoformat(), product, units, PRODUCTS[product]])
            total += units * PRODUCTS[product]
        _save_workbook(wb, out / f"sales_{region.lower()}.xlsx")
        revenue[region] = _money(total)
    return {"revenue": revenue, "total": _money(sum(revenue.values())), "best": max(revenue, key=revenue.get)}


def level_6(rng: random.Random, out: Path) -> Dict:
    """sales_by_rep.xlsx + rates.csv, where one rep has no commission rate: it must be reported, not invented."""
    reps = ["Aaron Blake", "Beatriz Silva", "Dana Whitfield", "Ethan Cole", "Fatima Zahra", "Liam O'Connor"]
    missing = "Dana Whitfield"
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Sales"
    sheet.append(["rep", "region", "sales"])
    sales = {}
    for rep in reps:
        sales[rep] = _money(rng.uniform(20000, 120000))
        sheet.append([rep, rng.choice(["North", "South", "East", "West"]), sales[rep]])
    _save_workbook(wb, out / "sales_by_rep.xlsx")
    rates = {rep: rng.choice([0.04, 0.05, 0.06, 0.075]) for rep in reps if rep != missing}
    _write_csv(out / "rates.csv", ["rep", "commission_rate"], [[rep, rate] for rep, rate in rates.items()])
    return {"sales": sales, "rates": rates, "missing": missing,
            "commission": {rep: _money(sales[rep] * rate) for rep, rate in rates.items()}}


LEVELS = {
    "office_1_sum_column": level_1,
    "office_2_xlsx_total": level_2,
    "office_3_clean_pivot": level_3,
    "office_4_mail_merge": level_4,
    "office_5_sales_report": level_5,
    "office_6_missing_rate": level_6,
}


def main() -> None:
    for task_id, build in LEVELS.items():
        out = FIXTURES_DIR / task_id
        shutil.rmtree(out, ignore_errors=True)
        out.mkdir(parents=True)
        expected = build(random.Random(f"{SEED}-{task_id}"), out)
        EXPECTED_DIR.mkdir(exist_ok=True)
        (EXPECTED_DIR / f"{task_id}.json").write_text(json.dumps(expected, indent=2) + "\n", encoding="utf-8")
        print(f"{task_id}: {', '.join(sorted(p.name for p in out.iterdir()))}")


if __name__ == "__main__":
    main()
