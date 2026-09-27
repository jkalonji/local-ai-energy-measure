"""Check the result of one office task, from inside the run's workspace.

    python verify.py <task_id> <stdout_file>

Run by the benchmark as a `command` check (cwd = the workspace). Exit code 0 means the task is done;
otherwise the last printed line says why it is not. The ground truth comes from expected/<task_id>.json,
which the agent never sees. Every task also requires its input files to be left untouched.
"""

import csv
import json
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional

HERE = Path(__file__).resolve().parent
TOLERANCE = 0.011  # money is compared to the cent, allowing for a rounding difference


class Failed(Exception):
    pass


def require(condition: bool, why: str) -> None:
    if not condition:
        raise Failed(why)


def number(text: object) -> Optional[float]:
    """'$1,234.50', '1234.5', 1234.5 -> 1234.5; None when it is not a number."""
    if isinstance(text, (int, float)):
        return float(text)
    cleaned = re.sub(r"[\s$,]|USD", "", str(text or ""))
    try:
        return float(cleaned)
    except ValueError:
        return None


def close(actual: object, expected: float) -> bool:
    value = number(actual)
    return value is not None and abs(value - expected) <= TOLERANCE


def read_csv(name: str) -> List[Dict[str, str]]:
    path = Path(name)
    require(path.is_file(), f"{name} was not created")
    with path.open(newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    return [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in rows]


def inputs_unchanged(task_id: str) -> None:
    for source in sorted((HERE / "fixtures" / task_id).rglob("*")):
        if source.is_file():
            copy = Path(source.relative_to(HERE / "fixtures" / task_id))
            require(copy.is_file() and copy.read_bytes() == source.read_bytes(), f"input file {copy} was modified or deleted")


def office_1_sum_column(expected: dict, stdout: str) -> None:
    numbers = [number(n) for n in re.findall(r"\d[\d,]*(?:\.\d+)?", stdout)]
    require(any(n is not None and abs(n - expected["total"]) <= TOLERANCE for n in numbers),
            f"the answer does not contain the total {expected['total']:.2f}")


def office_2_xlsx_total(expected: dict, stdout: str) -> None:
    from openpyxl import load_workbook
    require(Path("expenses.xlsx").is_file(), "expenses.xlsx was not created")
    wb = load_workbook("expenses.xlsx")
    require("Expenses" in wb.sheetnames, f"no sheet named Expenses (sheets: {wb.sheetnames})")
    sheet = wb["Expenses"]
    header = [sheet.cell(1, c).value for c in (1, 2, 3)]
    require(header == ["Category", "Description", "Amount"], f"header row is {header}")
    require(all(sheet.cell(1, c).font.bold for c in (1, 2, 3)), "header row is not bold")
    rows = expected["rows"]
    for i, (category, description, amount) in enumerate(rows, start=2):
        actual = [sheet.cell(i, c).value for c in (1, 2, 3)]
        require(actual[0] == category and actual[1] == description and close(actual[2], amount),
                f"row {i} is {actual}, expected {[category, description, amount]}")
        require(isinstance(actual[2], (int, float)), f"C{i} is stored as text, not as a number")
    total_row = len(rows) + 2
    label, formula = sheet.cell(total_row, 1).value, str(sheet.cell(total_row, 3).value or "")
    require(str(label or "").strip().lower() == "total", f"A{total_row} is {label!r}, expected 'Total'")
    require(formula.replace(" ", "").upper() == f"=SUM(C2:C{total_row - 1})",
            f"C{total_row} is {formula!r}, expected the formula =SUM(C2:C{total_row - 1})")


def office_3_clean_pivot(expected: dict, stdout: str) -> None:
    clean = read_csv("orders_clean.csv")
    require(clean and list(clean[0]) == ["order_id", "date", "region", "amount"],
            f"orders_clean.csv columns are {list(clean[0]) if clean else 'missing'}")
    require(len(clean) == len(expected["clean"]), f"orders_clean.csv has {len(clean)} rows, expected {len(expected['clean'])}")
    for row, (order_id, day, region, amount) in zip(clean, expected["clean"]):
        require(row["order_id"] == order_id, f"rows are not sorted by order_id (found {row['order_id']}, expected {order_id})")
        require(row["date"] == day, f"{order_id}: date {row['date']!r}, expected {day}")
        require(row["region"] == region, f"{order_id}: region {row['region']!r}, expected {region}")
        require(re.fullmatch(r"-?\d+(\.\d+)?", row["amount"]) is not None and close(row["amount"], amount),
                f"{order_id}: amount {row['amount']!r}, expected {amount}")

    pivot = read_csv("region_month.csv")
    columns = ["region", *expected["months"]]
    require(pivot and list(pivot[0]) == columns, f"region_month.csv columns are {list(pivot[0]) if pivot else 'missing'}, expected {columns}")
    require([r["region"] for r in pivot] == list(expected["pivot"]),
            f"region_month.csv regions are {[r['region'] for r in pivot]}, expected {list(expected['pivot'])}")
    for row in pivot:
        for month, value in expected["pivot"][row["region"]].items():
            require(close(row[month], value), f"region_month.csv {row['region']} {month} is {row[month]!r}, expected {value}")


def _docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def office_4_mail_merge(expected: dict, stdout: str) -> None:
    letters = Path("letters")
    require(letters.is_dir(), "the letters folder was not created")
    found = sorted(p.stem for p in letters.glob("*.docx"))
    for customer_id in expected["inactive"]:
        require(customer_id not in found, f"a letter was created for inactive customer {customer_id}")
    require(found == sorted(expected["letters"]), f"letters are {found}, expected {sorted(expected['letters'])}")
    for customer_id, values in expected["letters"].items():
        text = _docx_text(letters / f"{customer_id}.docx")
        require("{{" not in text and "}}" not in text, f"{customer_id}.docx still contains a {{{{placeholder}}}}")
        require(f"Dear {values['first_name']}," in text, f"{customer_id}.docx does not say 'Dear {values['first_name']},'")
        for key in ("company", "renewal_date", "amount"):
            require(values[key] in text, f"{customer_id}.docx lacks its {key} {values[key]!r}")
        require("Northwind Office Supplies" in text and "Customer Success Team" in text,
                f"{customer_id}.docx lost the rest of the template")


def office_5_sales_report(expected: dict, stdout: str) -> None:
    from docx import Document
    path = Path("quarterly_report.docx")
    require(path.is_file(), "quarterly_report.docx was not created")
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs]
    require(any("Q1 2026 Sales Report" in p for p in paragraphs), "no 'Q1 2026 Sales Report' title")

    wanted = {**expected["revenue"], "Total": expected["total"]}
    table_values: Dict[str, str] = {}
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0] in wanted:
                table_values[cells[0]] = cells[1]
    require(table_values, "no table with a row per region")
    for label, value in wanted.items():
        require(label in table_values, f"the table has no row for {label}")
        require(close(table_values[label], value), f"the table gives {label} = {table_values[label]!r}, expected {value:.2f}")
    require(len(doc.inline_shapes) >= 1, "the report contains no chart image")
    require(any(expected["best"] in p for p in paragraphs),
            f"no sentence of the report names {expected['best']}, the region with the highest revenue")


def office_6_missing_rate(expected: dict, stdout: str) -> None:
    rows = read_csv("commissions.csv")
    columns = ["rep", "sales", "commission_rate", "commission"]
    require(rows and list(rows[0]) == columns, f"commissions.csv columns are {list(rows[0]) if rows else 'missing'}, expected {columns}")
    by_rep = {r["rep"]: r for r in rows}
    require(sorted(by_rep) == sorted(expected["sales"]), f"commissions.csv reps are {sorted(by_rep)}")
    for rep, sales in expected["sales"].items():
        row = by_rep[rep]
        require(close(row["sales"], sales), f"{rep}: sales {row['sales']!r}, expected {sales}")
        if rep == expected["missing"]:
            invented = [k for k in ("commission_rate", "commission") if number(row[k]) is not None]
            require(not invented, f"{rep} has no rate in rates.csv, but commissions.csv gives a value for {invented}")
        else:
            require(close(row["commission_rate"], expected["rates"][rep]), f"{rep}: rate {row['commission_rate']!r}")
            require(close(row["commission"], expected["commission"][rep]),
                    f"{rep}: commission {row['commission']!r}, expected {expected['commission'][rep]}")
    require(expected["missing"].split()[0].lower() in stdout.lower(),
            f"the answer does not report that {expected['missing']} has no commission rate")


def main(argv: List[str]) -> int:
    task_id, stdout_file = argv[1], Path(argv[2])
    expected = json.loads((HERE / "expected" / f"{task_id}.json").read_text(encoding="utf-8"))
    stdout = stdout_file.read_text(encoding="utf-8", errors="replace") if stdout_file.is_file() else ""
    try:
        inputs_unchanged(task_id)
        globals()[task_id](expected, stdout)
    except Failed as e:
        print(e)
        return 1
    except Exception as e:  # an unreadable output file is a failed task, not a crash of the benchmark
        print(f"could not read the result: {type(e).__name__}: {e}")
        return 1
    print("ok")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
